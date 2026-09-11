#!/usr/bin/env python3
import argparse
import base64
import hashlib
import os
import re
from datetime import datetime
from pathlib import Path

try:
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit(
        "Dependencias ausentes. Instale no servidor com:\n"
        "python3 -m pip install --user openpyxl cryptography\n"
        f"Detalhe: {exc}"
    )


DEFAULT_INPUT = "Planilha de Senhas (Oficinas) - atualizado em 25.06.xlsx"
DEFAULT_OUTPUT = "outputs/import_portais_senhas/import_portais_senhas.sql"
REQUIRED_HEADERS = [
    "SEGMENTO",
    "SEGURADORA",
    "EMPRESA",
    "ACESSO",
    "SENHA",
    "idEmpresa",
    "idSeguradora",
]

SEGURADORA_DEPARA = {
    "ALLIANZ (AR)": "Allianz",
    "ALLIANZ (FRN Direto)": "Allianz",
    "AZUL (Ofc)": "Azul",
    "ESSOR (Ofc)": "Essor",
    "GENERALI (Ofc)": "Generali",
    "GENTE (Ofc)": "Gente",
    "HDI (Ofc)": "HDI",
    "MAPFRE (AR Atri e Jeep)": "Mapfre",
    "PORTO SEGURO (OFC)": "Porto",
    "SANTANDER (Ofc)": "Santander Auto",
    "SOMPO (AR Audi RP / AON)": "Sompo",
    "SOMPO (AR)": "Sompo",
    "SURA (AR Ortovel, Audi e Thor)": "Sura",
    "TOKIO (AR Euro / Keiji)": "Tokio",
    "TOKIO MARINE (Ofc)": "Tokio",
    "YELUM (Ofc)": "Yelum",
}


def load_env(path):
    values = {}
    if not path.exists():
        return values

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
            value = value[1:-1]
        values[key.strip()] = value
    return values


def sql_string(value):
    if value is None:
        return "NULL"
    text = str(value)
    return "'" + text.replace("\\", "\\\\").replace("'", "''") + "'"


def normalize_cell(value):
    return str(value).strip() if value is not None else ""


def normalize_cnpj(value):
    return re.sub(r"\D+", "", normalize_cell(value))[:14]


def truncate(value, length):
    value = normalize_cell(value)
    return value[:length] if value else ""


def looks_like_url(value):
    text = normalize_cell(value).lower()
    return text.startswith(("http://", "https://")) or text.startswith("www.") or bool(re.search(r"\.[a-z]{2,}(/|$)", text))


def encrypt_password(secret, password):
    key = hashlib.sha256(secret.encode("utf-8")).digest()
    iv = os.urandom(12)
    encrypted_with_tag = AESGCM(key).encrypt(iv, password.encode("utf-8"), None)
    encrypted = encrypted_with_tag[:-16]
    tag = encrypted_with_tag[-16:]
    return {
        "encrypted": base64.b64encode(encrypted).decode("ascii"),
        "iv": base64.b64encode(iv).decode("ascii"),
        "tag": base64.b64encode(tag).decode("ascii"),
    }


HEADER_ALIASES = {
    "ACESSO": ["ACESSO", "ACESSO (Cobrança)", "ACESSO (Cobranca)"],
    "USUARIO": ["USUÁRIO2", "USUARIO2", "USUÁRIO", "USUARIO"],
    "idEmpresa": ["idEmpresa", "IDEMPRESA", "ID EMPRESA", "empresaId", "EMPRESA ID"],
    "idSeguradora": ["idSeguradora", "IDSEGURADORA", "ID SEGURADORA", "seguradoraId", "SEGURADORA ID"],
}


def resolve_header(headers, header):
    candidates = HEADER_ALIASES.get(header, [header])
    normalized = {normalize_header(value): value for value in headers}
    for candidate in candidates:
        found = normalized.get(normalize_header(candidate))
        if found:
            return found
    return None


def normalize_header(value):
    return re.sub(r"[^A-Z0-9]+", "", normalize_cell(value).upper())


def get_value(data, header):
    resolved = resolve_header(data.keys(), header)
    return data.get(resolved, "") if resolved else ""


def read_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = [normalize_cell(cell.value) for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    missing = [header for header in REQUIRED_HEADERS if not resolve_header(headers, header)]
    if missing:
        raise RuntimeError(f"Cabecalhos ausentes na planilha: {', '.join(missing)}")

    rows = []
    skipped = []
    for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        data = {header: normalize_cell(cells[index] if index < len(cells) else "") for index, header in enumerate(headers)}
        if not any(data.values()):
            continue

        missing_fields = [
            label
            for label in ["ACESSO", "SENHA", "idEmpresa", "idSeguradora"]
            if not get_value(data, label)
        ]
        if missing_fields:
            skipped.append((row_number, missing_fields))
            continue
        invalid_ids = [
            label
            for label in ["idEmpresa", "idSeguradora"]
            if not get_value(data, label).isdigit()
        ]
        if invalid_ids:
            skipped.append((row_number, [f"{label} invalido" for label in invalid_ids]))
            continue

        rows.append((row_number, data))
    return rows, skipped


def build_observacao(data, row_number):
    parts = [
        f"Origem: planilha de senhas de oficinas, linha {row_number}",
        f"Segmento: {data.get('SEGMENTO', '')}",
        f"Seguradora: {data.get('SEGURADORA', '')}",
        f"Empresa: {data.get('EMPRESA', '')}",
        f"CNPJ: {data.get('CNPJ', '')}",
        f"Acesso original: {get_value(data, 'ACESSO')}",
    ]
    return " | ".join(part for part in parts if not part.endswith(": "))


def generate_sql(input_path, output_path, database, secret):
    rows, skipped = read_rows(input_path)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    statements = [
        "-- Importacao de senhas de portais gerada automaticamente.",
        f"-- Gerado em: {now}",
        f"-- Origem: {input_path}",
        f"-- Linhas validas para carga: {len(rows)}",
        f"-- Linhas ignoradas por campos obrigatorios ausentes: {len(skipped)}",
        "SET NAMES utf8mb4;",
        f"USE `{database}`;",
        "START TRANSACTION;",
        "",
    ]

    if skipped:
        statements.append("-- Linhas ignoradas:")
        for row_number, missing_fields in skipped:
            statements.append(f"-- linha {row_number}: ausente {', '.join(missing_fields)}")
        statements.append("")

    for row_number, data in rows:
        segmento = truncate(data["SEGMENTO"], 80)
        seguradora_planilha = truncate(data["SEGURADORA"], 180)
        empresa_nome = truncate(data["EMPRESA"], 180)
        cnpj = normalize_cnpj(data.get("CNPJ", ""))
        empresa_id = int(get_value(data, "idEmpresa"))
        seguradora_id = int(get_value(data, "idSeguradora"))
        portal_nome = truncate(get_value(data, "ACESSO"), 120)
        acesso = truncate(get_value(data, "ACESSO"), 255)
        portal_url = acesso if looks_like_url(acesso) else None
        usuario = truncate(get_value(data, "USUARIO") or acesso, 180)
        observacao = truncate(build_observacao(data, row_number), 1000)
        encrypted = encrypt_password(secret, data["SENHA"])

        statements.extend(
            [
                f"-- Linha {row_number}",
                f"SET @empresaId := {empresa_id};",
                f"SET @seguradoraId := {seguradora_id};",
                "SET @portalCredentialId := (",
                "  SELECT id",
                "  FROM portais_senhas",
                "  WHERE empresaId = @empresaId",
                "    AND seguradoraId = @seguradoraId",
                f"    AND portalNome = {sql_string(portal_nome)}",
                f"    AND usuario = {sql_string(usuario)}",
                "  ORDER BY id",
                "  LIMIT 1",
                ");",
                "UPDATE portais_senhas",
                "SET",
                "  empresaId = @empresaId,",
                "  seguradoraId = @seguradoraId,",
                f"  segmento = {sql_string(segmento)},",
                f"  empresaNome = {sql_string(empresa_nome)},",
                f"  cnpj = {sql_string(cnpj or None)},",
                f"  portalUrl = {sql_string(portal_url)},",
                f"  senhaCriptografada = {sql_string(encrypted['encrypted'])},",
                f"  senhaIv = {sql_string(encrypted['iv'])},",
                f"  senhaTag = {sql_string(encrypted['tag'])},",
                f"  observacao = {sql_string(observacao)},",
                "  ativo = 1,",
                "  atualizadoEm = CURRENT_TIMESTAMP",
                "WHERE id = @portalCredentialId;",
                "INSERT INTO portais_senhas (",
                "  empresaId, seguradoraId, segmento, empresaNome, cnpj,",
                "  portalNome, portalUrl, usuario,",
                "  senhaCriptografada, senhaIv, senhaTag, observacao, ativo",
                ")",
                "SELECT",
                f"  @empresaId, @seguradoraId, {sql_string(segmento)}, {sql_string(empresa_nome)}, {sql_string(cnpj or None)},",
                f"  {sql_string(portal_nome)}, {sql_string(portal_url)}, {sql_string(usuario)},",
                f"  {sql_string(encrypted['encrypted'])}, {sql_string(encrypted['iv'])}, {sql_string(encrypted['tag'])}, {sql_string(observacao)}, 1",
                "WHERE @portalCredentialId IS NULL;",
                "",
            ]
        )

    statements.extend(
        [
            "COMMIT;",
            "",
            "SELECT",
            f"  {len(rows)} AS linhas_validas_processadas,",
            f"  {len(skipped)} AS linhas_ignoradas;",
        ]
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(statements) + "\n", encoding="utf-8")
    return len(rows), len(skipped)


def main():
    parser = argparse.ArgumentParser(description="Gera SQL de importacao para portais_senhas a partir da planilha de senhas.")
    parser.add_argument("--input", default=DEFAULT_INPUT)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    parser.add_argument("--database", default="rpa_oficinas")
    parser.add_argument("--env", default=".env")
    args = parser.parse_args()

    env_values = load_env(Path(args.env))
    secret = (
        os.getenv("PORTAL_CREDENTIALS_SECRET")
        or env_values.get("PORTAL_CREDENTIALS_SECRET")
        or os.getenv("APP_SECRET")
        or env_values.get("APP_SECRET")
        or os.getenv("NEXTAUTH_SECRET")
        or env_values.get("NEXTAUTH_SECRET")
    )
    if not secret or len(secret) < 16:
        raise RuntimeError("Configure PORTAL_CREDENTIALS_SECRET, APP_SECRET ou NEXTAUTH_SECRET com pelo menos 16 caracteres.")

    inserted, skipped = generate_sql(Path(args.input), Path(args.output), args.database, secret)
    print(f"SQL gerado em {args.output}")
    print(f"Linhas validas: {inserted}")
    print(f"Linhas ignoradas: {skipped}")


if __name__ == "__main__":
    main()
